from pathlib import Path
import csv

from pptx import Presentation
from pptx.chart.data import CategoryChartData, XyChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


BASE_DIR = Path(__file__).resolve().parent
ASSETS_DIR = BASE_DIR / "ppt-assets"
OUTPUT_FILE = BASE_DIR / "AlcoveHire-Investor-Deck-v1.pptx"


PRIMARY = RGBColor(0x0F, 0x4C, 0x81)
ACCENT = RGBColor(0x1E, 0x7A, 0x73)
TEXT_DARK = RGBColor(0x1D, 0x27, 0x36)
MUTED = RGBColor(0x5E, 0x6B, 0x7E)
BG_LIGHT = RGBColor(0xF4, 0xF7, 0xFB)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)


def set_bg(slide):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = BG_LIGHT


def style_title(shape):
    tf = shape.text_frame
    p = tf.paragraphs[0]
    if not p.runs:
        p.add_run()
    run = p.runs[0]
    run.font.name = "Manrope"
    run.font.bold = True
    run.font.size = Pt(34)
    run.font.color.rgb = PRIMARY


def add_header(slide, title):
    title_box = slide.shapes.add_textbox(Inches(0.6), Inches(0.3), Inches(12.0), Inches(0.8))
    title_box.text_frame.text = title
    style_title(title_box)
    return title_box


def add_bullets(slide, x, y, w, h, lines, font_size=20):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    for idx, line in enumerate(lines):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.text = line
        p.level = 0
        p.space_after = Pt(8)
        p.font.name = "Source Sans 3"
        p.font.size = Pt(font_size)
        p.font.color.rgb = TEXT_DARK
    return box


def add_footer(slide, page_no):
    left = slide.shapes.add_textbox(Inches(0.6), Inches(7.0), Inches(4.0), Inches(0.3))
    left.text_frame.text = "AlcoveHire"
    p = left.text_frame.paragraphs[0]
    p.font.name = "Source Sans 3"
    p.font.size = Pt(11)
    p.font.color.rgb = MUTED

    right = slide.shapes.add_textbox(Inches(9.5), Inches(7.0), Inches(3.2), Inches(0.3))
    right.text_frame.text = f"Confidential | Jul 2026 | {page_no}"
    p2 = right.text_frame.paragraphs[0]
    p2.alignment = PP_ALIGN.RIGHT
    p2.font.name = "Source Sans 3"
    p2.font.size = Pt(11)
    p2.font.color.rgb = MUTED


def read_csv_dict(path):
    with path.open("r", encoding="utf-8") as f:
        return list(csv.DictReader(f))


def add_chart_title(slide, text, x, y, w, h=0.3):
    t = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    t.text_frame.text = text
    p = t.text_frame.paragraphs[0]
    p.font.name = "Source Sans 3"
    p.font.bold = True
    p.font.size = Pt(14)
    p.font.color.rgb = PRIMARY


def build():
    prs = Presentation()

    # Slide 1: Cover
    s = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(s)
    box = s.shapes.add_textbox(Inches(0.8), Inches(1.3), Inches(6.5), Inches(2.5))
    tf = box.text_frame
    tf.text = "AlcoveHire"
    p = tf.paragraphs[0]
    p.font.name = "Manrope"
    p.font.bold = True
    p.font.size = Pt(54)
    p.font.color.rgb = PRIMARY

    p2 = tf.add_paragraph()
    p2.text = "AI-Powered Recruitment Platform for High-Velocity Hiring Teams"
    p2.font.name = "Source Sans 3"
    p2.font.size = Pt(24)
    p2.font.color.rgb = TEXT_DARK
    p2.space_before = Pt(16)

    p3 = tf.add_paragraph()
    p3.text = "Alcove Infotech | Investor Deck"
    p3.font.name = "Source Sans 3"
    p3.font.size = Pt(16)
    p3.font.color.rgb = MUTED
    p3.space_before = Pt(20)

    # decorative card
    card = s.shapes.add_shape(1, Inches(8.0), Inches(1.2), Inches(4.5), Inches(4.8))
    card.fill.solid()
    card.fill.fore_color.rgb = WHITE
    card.line.color.rgb = RGBColor(0xD8, 0xE0, 0xEA)
    txt = card.text_frame
    txt.text = "Vision"
    rp = txt.paragraphs[0]
    rp.font.name = "Manrope"
    rp.font.bold = True
    rp.font.size = Pt(20)
    rp.font.color.rgb = PRIMARY
    rp.alignment = PP_ALIGN.CENTER
    rp2 = txt.add_paragraph()
    rp2.text = "Faster shortlists\nBetter conversions\nPredictable recruiter productivity"
    rp2.font.name = "Source Sans 3"
    rp2.font.size = Pt(16)
    rp2.font.color.rgb = TEXT_DARK
    rp2.alignment = PP_ALIGN.CENTER

    # Slide 2: Problem + bar chart
    s = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(s)
    add_header(s, "Problem: Fragmented Recruiting Workflows Hurt Placements")
    add_bullets(
        s,
        0.7,
        1.3,
        6.0,
        4.8,
        [
            "Candidate data scattered across sheets, email, and chat.",
            "Slow requirement-to-shortlist turnaround.",
            "Limited visibility into team productivity and pipeline leak.",
            "Legacy ATS tools are costly for SMB teams.",
        ],
        font_size=19,
    )

    rows = read_csv_dict(ASSETS_DIR / "problem-time-loss.csv")
    chart_data = CategoryChartData()
    chart_data.categories = [r["Activity"].replace("_", " ") for r in rows]
    chart_data.add_series("Hours", [float(r["HoursLostPerWeek"]) for r in rows])
    chart = s.shapes.add_chart(
        XL_CHART_TYPE.BAR_CLUSTERED,
        Inches(7.0), Inches(1.5), Inches(5.5), Inches(4.3),
        chart_data,
    ).chart
    chart.has_legend = False
    chart.value_axis.has_major_gridlines = True
    add_chart_title(s, "Weekly Time Lost (Hours)", 7.0, 1.15, 5.0)
    add_footer(s, 2)

    # Slide 3: Solution
    s = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(s)
    add_header(s, "Solution: One AI-Powered Operating Platform")
    add_bullets(
        s,
        0.8,
        1.4,
        11.8,
        5.2,
        [
            "Unified ATS workflows: candidates, clients, requirements, interviews, invoices.",
            "Embedded AI for parsing, scoring, summarization, and outreach drafting.",
            "Role-based dashboards for recruiters, managers, and owners.",
            "Credit-based AI usage model to control cost and grow ARPU.",
        ],
        font_size=22,
    )
    add_footer(s, 3)

    # Slide 4: Product flow
    s = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(s)
    add_header(s, "Product Flow: Requirement to Placement")
    flow_steps = ["Requirement", "AI Match", "Interview", "Offer", "Invoice"]
    start_x = 0.8
    for i, step in enumerate(flow_steps):
        shape = s.shapes.add_shape(1, Inches(start_x + i * 2.35), Inches(2.6), Inches(2.1), Inches(1.2))
        shape.fill.solid()
        shape.fill.fore_color.rgb = WHITE
        shape.line.color.rgb = PRIMARY
        tf = shape.text_frame
        tf.text = step
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        p.font.name = "Manrope"
        p.font.bold = True
        p.font.size = Pt(18)
        p.font.color.rgb = PRIMARY
    add_bullets(
        s,
        0.8,
        4.3,
        12.0,
        2.0,
        [
            "Result: faster shortlist cycles, better candidate quality, and billing visibility.",
        ],
        font_size=20,
    )
    add_footer(s, 4)

    # Slide 5: Market opportunity
    s = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(s)
    add_header(s, "Market Opportunity: TAM / SAM / SOM")
    market = read_csv_dict(ASSETS_DIR / "market-tam-sam-som.csv")
    cdata = CategoryChartData()
    cdata.categories = [r["Segment"] for r in market]
    cdata.add_series("INR Cr", [float(r["ValueINRCr"]) for r in market])
    chart = s.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED,
        Inches(0.9), Inches(1.6), Inches(7.0), Inches(4.6),
        cdata,
    ).chart
    chart.has_legend = False
    add_bullets(
        s,
        8.1,
        1.8,
        4.2,
        4.8,
        [
            "TAM: INR 8,000 Cr",
            "SAM: INR 1,800 Cr",
            "SOM (3Y): INR 180 Cr",
            "Focused entry via agency-first motion",
        ],
        font_size=18,
    )
    add_footer(s, 5)

    # Slide 6: Business model and pricing
    s = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(s)
    add_header(s, "Business Model and Pricing")
    add_bullets(
        s,
        0.8,
        1.5,
        6.5,
        3.0,
        [
            "Starter: INR 2,999 / month",
            "Growth: INR 7,999 / month",
            "Business: INR 18,999 / month",
            "Yearly billing discount: 20%",
        ],
        font_size=19,
    )
    mix = read_csv_dict(ASSETS_DIR / "revenue-mix.csv")
    pdata = CategoryChartData()
    pdata.categories = [r["RevenueStream"].replace("_", " ") for r in mix]
    pdata.add_series("Share", [float(r["Percent"]) for r in mix])
    pie = s.shapes.add_chart(
        XL_CHART_TYPE.PIE,
        Inches(7.1), Inches(1.6), Inches(5.0), Inches(4.5),
        pdata,
    ).chart
    pie.has_legend = True
    pie.legend.position = XL_LEGEND_POSITION.BOTTOM
    add_chart_title(s, "Revenue Mix", 8.4, 1.25, 3.0)
    add_footer(s, 6)

    # Slide 7: GTM funnel
    s = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(s)
    add_header(s, "Go-To-Market Funnel")
    gtm = read_csv_dict(ASSETS_DIR / "gtm-funnel.csv")
    fdata = CategoryChartData()
    fdata.categories = [r["Stage"].replace("_", " ") for r in gtm]
    fdata.add_series("Count", [float(r["Count"]) for r in gtm])
    chart = s.shapes.add_chart(
        XL_CHART_TYPE.BAR_CLUSTERED,
        Inches(0.9), Inches(1.6), Inches(7.6), Inches(4.8),
        fdata,
    ).chart
    chart.has_legend = False
    add_bullets(
        s,
        8.7,
        1.8,
        3.8,
        4.8,
        [
            "Founder-led sales for first 20 accounts",
            "Webinars + outbound for next 80 accounts",
            "14-day trial to paid conversion model",
        ],
        font_size=17,
    )
    add_footer(s, 7)

    # Slide 8: Competitive positioning (scatter)
    s = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(s)
    add_header(s, "Competitive Positioning")
    cmp_rows = read_csv_dict(ASSETS_DIR / "competitive-quadrant.csv")
    xy_data = XyChartData()
    for row in cmp_rows:
        series = xy_data.add_series(row["Company"].replace("_", " "))
        series.add_data_point(float(row["WorkflowDepth"]), float(row["PriceToValue"]))
    chart = s.shapes.add_chart(
        XL_CHART_TYPE.XY_SCATTER,
        Inches(0.9), Inches(1.6), Inches(7.8), Inches(4.8),
        xy_data,
    ).chart
    chart.has_legend = True
    chart.legend.position = XL_LEGEND_POSITION.RIGHT
    add_bullets(
        s,
        8.9,
        1.8,
        3.6,
        4.8,
        [
            "High workflow depth",
            "Strong SMB price-to-value fit",
            "AI built into daily recruiter workflow",
        ],
        font_size=17,
    )
    add_footer(s, 8)

    # Slide 9: Milestones
    s = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(s)
    add_header(s, "Execution Roadmap and Milestones")
    milestones = read_csv_dict(ASSETS_DIR / "milestones.csv")
    y = 1.7
    for row in milestones:
        chip = s.shapes.add_shape(1, Inches(0.9), Inches(y), Inches(2.0), Inches(0.65))
        chip.fill.solid()
        chip.fill.fore_color.rgb = ACCENT
        chip.line.color.rgb = ACCENT
        ctf = chip.text_frame
        ctf.text = row["Month"]
        cp = ctf.paragraphs[0]
        cp.alignment = PP_ALIGN.CENTER
        cp.font.name = "Manrope"
        cp.font.bold = True
        cp.font.size = Pt(16)
        cp.font.color.rgb = WHITE

        box = s.shapes.add_shape(1, Inches(3.1), Inches(y), Inches(8.8), Inches(0.65))
        box.fill.solid()
        box.fill.fore_color.rgb = WHITE
        box.line.color.rgb = RGBColor(0xD5, 0xDE, 0xE8)
        tf = box.text_frame
        tf.text = row["Milestone"]
        p = tf.paragraphs[0]
        p.font.name = "Source Sans 3"
        p.font.size = Pt(16)
        p.font.color.rgb = TEXT_DARK
        y += 0.85
    add_footer(s, 9)

    # Slide 10: Financial snapshot
    s = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(s)
    add_header(s, "Financial Snapshot: MRR vs Burn")
    fin = read_csv_dict(ASSETS_DIR / "mrr-vs-burn.csv")
    line = CategoryChartData()
    line.categories = [r["Month"] for r in fin]
    line.add_series("MRR (INR Lakhs)", [float(r["MRR_Lakhs"]) for r in fin])
    line.add_series("Burn (INR Lakhs)", [float(r["Burn_Lakhs"]) for r in fin])
    ch = s.shapes.add_chart(
        XL_CHART_TYPE.LINE_MARKERS,
        Inches(0.9), Inches(1.6), Inches(7.8), Inches(4.8),
        line,
    ).chart
    ch.has_legend = True
    ch.legend.position = XL_LEGEND_POSITION.BOTTOM
    add_bullets(
        s,
        8.9,
        1.8,
        3.6,
        4.8,
        [
            "MRR at 100 customers: INR 6.1L",
            "ARR at 100 customers: INR 73.2L",
            "Burn at full ramp: INR 17.9L/month",
        ],
        font_size=17,
    )
    add_footer(s, 10)

    # Slide 11: Funding ask
    s = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(s)
    add_header(s, "Funding Ask: INR 2.2 Cr")
    use = read_csv_dict(ASSETS_DIR / "use-of-funds.csv")
    udata = CategoryChartData()
    udata.categories = [r["Category"].replace("_", " ") for r in use]
    udata.add_series("Percent", [float(r["Percent"]) for r in use])
    doughnut = s.shapes.add_chart(
        XL_CHART_TYPE.DOUGHNUT,
        Inches(0.9), Inches(1.6), Inches(7.4), Inches(4.8),
        udata,
    ).chart
    doughnut.has_legend = True
    doughnut.legend.position = XL_LEGEND_POSITION.RIGHT
    add_bullets(
        s,
        8.6,
        2.0,
        4.0,
        4.2,
        [
            "12-month runway",
            "Build, launch, and reach first 100 customers",
            "Milestone-based hiring and spend discipline",
        ],
        font_size=18,
    )
    add_footer(s, 11)

    # Slide 12: Team and close
    s = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(s)
    add_header(s, "Team and Closing")
    add_bullets(
        s,
        0.9,
        1.7,
        11.8,
        4.8,
        [
            "Core team: 2 Frontend, 2 Backend, 1 QA, 1 DevOps/Security, 1 PM, 1 UI/UX",
            "Go-to-market: 1 Sales Lead, 2 Sales Executives, 1 Customer Success, 1 Marketer",
            "Delivery commitment: launch-ready product in 4 months",
            "Ask: proceed to due diligence and term-sheet discussion",
        ],
        font_size=20,
    )
    add_footer(s, 12)

    prs.save(OUTPUT_FILE)
    print(f"Created: {OUTPUT_FILE}")


if __name__ == "__main__":
    build()
